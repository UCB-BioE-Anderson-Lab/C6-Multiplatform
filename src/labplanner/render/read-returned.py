#!/usr/bin/env python3
"""Read a filled-in labsheet back — what the student wrote, by slug.

    read-returned.py <filled.xlsx> [--json]

**The record tab is a map, not a cache.** Each row is `slug` and a formula `='Sheet'!A16`, and
this resolves the reference itself rather than reading the formula's cached value. That matters
more than it looks: a cached value exists only if a real spreadsheet application computed it, so
trusting the cache makes extraction depend on which program the student happened to use — and
fail silently, with empty fields, when they used one that did not write caches back.

Following the reference works whatever wrote the file, and it is the same information.

**A blank field comes back as blank, never as absent.** A student who did not fill something in
and a slug that no longer exists are different facts, and a reader that conflates them will
report an empty inventory as confidently as a full one.
"""
import json, re, sys
import openpyxl

REF = re.compile(r"^=\s*'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)\s*$")


# The hidden tab the renderer writes, defaulted to the neutral name it uses. A lab that renames
# it (Cortex calls it `cortex-record`) passes the name back in here; nothing about reading a
# returned workbook is that lab's business.
RECORD_TAB = "record"

def read_returned(path, record_tab=None):
    """{slug: value} for every field the sheet declares. -> (values, problems)"""
    wb = openpyxl.load_workbook(path, data_only=False)
    problems = []
    tab = record_tab or RECORD_TAB
    if tab not in wb.sheetnames:
        return {}, [f"no {tab} tab — this file was not made by labpacket-to-xlsx, "
                    "or the tab was deleted"]
    out = {}
    for row in wb[tab].iter_rows(min_row=2, max_col=2, values_only=True):
        slug, formula = (row + (None, None))[:2]
        if not slug:
            continue
        m = REF.match(str(formula or ""))
        if not m:
            problems.append(f"{slug}: cannot read the reference {formula!r}")
            out[slug] = None
            continue
        sheet, col, rownum = m.group(1), m.group(2), int(m.group(3))
        if sheet not in wb.sheetnames:
            problems.append(f"{slug}: points at sheet {sheet!r}, which is not in this workbook")
            out[slug] = None
            continue
        v = wb[sheet][f"{col}{rownum}"].value
        out[slug] = None if v is None or str(v).strip() == "" else v
    return out, problems


def rows_of(values, prefix):
    """Collapse indexed slugs back into rows: miniprep.sample.2.well -> rows[1]['well']."""
    rows = {}
    for slug, v in values.items():
        if not slug.startswith(prefix + "."):
            continue
        rest = slug[len(prefix) + 1:].split(".")
        if len(rest) != 2 or not rest[0].isdigit():
            continue
        rows.setdefault(int(rest[0]), {})[rest[1]] = v
    # A row where every cell is blank was never filled in; returning it would put empty samples
    # into an inventory.
    return [rows[k] for k in sorted(rows) if any(x is not None for x in rows[k].values())]


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        sys.exit("  read-returned.py <filled.xlsx> [--json]")
    tab = None
    if "--record-tab" in sys.argv:
        tab = sys.argv[sys.argv.index("--record-tab") + 1]
    values, problems = read_returned(args[0], tab)
    if "--json" in sys.argv:
        # THE PREFIX IS DERIVED, NOT NAMED. This said `rows_of(values, "miniprep.sample")`, and
        # the slugs a sheet actually registers are `<sheet-id>.sample.<n>.<column>` —
        # `s7-miniprep-sequencing.sample.1.well`. So the literal matched nothing, `samples` came
        # back `[]` on every workbook, and a caller reading it concluded the student had filled in
        # nothing. Every prefix present is collapsed instead, keyed by sheet.
        prefixes = sorted({slug.rsplit(".sample.", 1)[0]
                           for slug in values if ".sample." in slug})
        print(json.dumps({"values": values,
                          "samples": {p: rows_of(values, f"{p}.sample") for p in prefixes},
                          "problems": problems}, indent=2, default=str))
        return 0
    for p in problems:
        print(f"  ! {p}")
    filled = {k: v for k, v in values.items() if v is not None}
    print(f"  {len(filled)} of {len(values)} field(s) filled in")
    for k, v in sorted(filled.items()):
        print(f"    {k:34} {v}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
