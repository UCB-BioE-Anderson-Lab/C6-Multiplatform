"""Convert one labsheet workbook into a LabPacket JSON instance.

Fitted to the grammar in C6/docs/LABSHEET-TRANSLATION.md, from a labsheet that worked. The
three kinds of tab are all carried: operation tabs become inputs/samples/recipe/notes, and the
rest becomes an ordered `blocks[]` that preserves what was on the page in the order it was on it.

FAITHFUL, NOT CLEVER. Where a tab does something this does not model — a conditional branch, a
plate layout — the content is kept as a block of the right kind and NOT reinterpreted. Losing
it would be worse, and pretending to understand it would be worse still.
"""
import openpyxl, json, os, re, sys

def cells(ws):
    out=[]
    for r in ws.iter_rows(values_only=True):
        c=[("" if x is None else str(x).strip()) for x in r]
        while c and not c[-1]: c.pop()
        out.append(c)
    return out

# A section marker may carry content on the same line — Zymo's is
# `source: Enzyme freezer pcr rack "to Zymo"`. Requiring a bare marker dropped that tab's
# whole samples table, and with it `elution_volume`, which appears nowhere else.
SECTION = re.compile(r"^(source|sources|samples|reaction|notes)\s*:\s*(.*)$", re.I)
STEPNUM = re.compile(r"^(\d+)\)?[.)]?$")

def convert_tab(ws, name, descriptions=None):
    rows = [r for r in cells(ws) if any(r)]
    if not rows: return None

    # THE FIRST CELL IS THE TITLE — UNLESS IT PLAINLY IS NOT.
    #
    # Taking rows[0][0] on faith produced three nonsense sheets out of fifty: a labsheet titled
    # `#REF!` (a broken formula left in the workbook), one titled `samples:` (the title row was
    # simply never written), and one titled `No Gibson reaction involved` (a remark sitting
    # above the title). All three would have printed onto a page a student carries, and the
    # underlying defect — a formula that lost its reference — would have been laundered into a
    # plausible-looking document. So: recognise the cases, fall back to the tab name, and SAY
    # SO, because the workbook is what actually needs fixing.
    head = rows[0][0]
    bad = (not head
           or head.startswith("#")                        # #REF!, #N/A, #VALUE!
           or SECTION.match(head)                         # the title row is missing entirely
           or "for experiment" not in head.lower() and len(rows[0]) > 1)
    if bad:
        print(f"     ! {name}: first row is {head!r}, not a title — using the tab name. "
              f"Fix the workbook.")
        head = f"{name} (untitled in the workbook)"
    sheet = {"id": None, "title": head, "operation": head.split()[0],
             "metadata": {}, "inputs": [], "samples": [], "outputs": [],
             "recipe": None, "notes": [], "blocks": []}
    # When the first row was rejected it is still CONTENT — `samples:` is a section marker, and
    # skipping it as though it were a title threw away the whole table under it. Only a row
    # actually used as the title is consumed.
    i, section = (0 if bad else 1), None
    pending_header = None
    while i < len(rows):
        c = rows[i]; first = c[0]
        low = first.lower()
        if low.startswith("protocol:"):
            sheet["metadata"]["module"] = (first.split(":",1)[1].strip() or (c[1] if len(c)>1 else "")); i+=1; continue
        if low.startswith("program:"):
            sheet["metadata"]["program"] = (first.split(":",1)[1].strip() or (c[1] if len(c)>1 else "")); i+=1; continue
        # A LONE `True`/`False` NEAR THE TOP IS THE STEP'S ON/OFF SWITCH.
        #
        # Assembly tabs carry one, because the workbook is a template holding both a Golden Gate
        # and a Gibson tab and an experiment uses one of them. SLIP4_singles' Gibson tab is
        # `False` and titled "No Gibson reaction involved" — a step that must not appear on a
        # printed labsheet, and must not be deleted from the record either, because "this
        # experiment deliberately did not do a Gibson" is a fact worth keeping. So it is carried
        # with `applies: false`, and the renderers leave it off the page.
        if len([x for x in c if x]) == 1 and first in ("True", "False", "TRUE", "FALSE"):
            sheet["applies"] = first.lower() == "true"; i += 1; continue
        if low.startswith("thermocycler"):
            sheet["metadata"]["thermocycler"] = True; i+=1; continue
        m = SECTION.match(first)
        if m:
            section = m.group(1).lower(); pending_header=None
            trailing = (m.group(2) or "").strip()
            if trailing:
                # e.g. `source: Enzyme freezer pcr rack "to Zymo"` — a location, not a table.
                sheet["metadata"].setdefault(section + "_note", trailing)
                section = None
            i+=1; continue
        if section in ("source","sources","samples") :
            # header row then data rows until blank/section
            if pending_header is None:
                pending_header = [x for x in c if x] or c
                i+=1; continue
            if len(c) == 1:
                # A single-cell row inside a table is the next heading, not a one-column
                # sample. Without this, "Protocol" became a sample whose reaction was
                # "Protocol" — a fabricated row in somebody's experimental record.
                section = None; pending_header = None; continue
            row = dict(zip(pending_header, c[:len(pending_header)]))
            if any(row.values()):
                sheet["inputs" if section.startswith("source") else "samples"].append(row)
            i+=1; continue
        if section == "reaction":
            if len(c)>=2 and re.match(r"^[\d.]+$", c[0]):
                comp = {"volume_uL": float(c[0]),
                        "name": next((x for x in c[1:] if x and x.lower()!="ul"), ""),
                        "code": c[-1] if c[-1].endswith("____") else ""}
                sheet["recipe"] = sheet["recipe"] or {"components": []}
                sheet["recipe"]["components"].append(comp); i+=1; continue
            section=None; continue
        if section == "notes":
            if c[0] == "*":
                sheet["notes"].append(" ".join(x for x in c[1:] if x))
            elif sheet["notes"]:
                # A wrapped note continues on the next row, and those rows carry a stray "0.0"
                # in the first column in the source workbook. Concatenating it put "0.0" into
                # the middle of a SAFETY instruction — "take the enzyme cooler out of the
                # freezer 0.0 when you are actively using it". Drop a leading bare number: it
                # is spreadsheet residue, never part of the sentence.
                tail = list(c)
                if tail and re.fullmatch(r"[\d.]+", tail[0].strip()):
                    tail = tail[1:]
                sheet["notes"][-1] += " " + " ".join(x for x in tail if x)
            i+=1; continue
        # ---- unstructured: kinds B and C -------------------------------------------------
        sm = STEPNUM.match(first)
        if sm:
            sheet["blocks"].append({"kind":"step","text":" ".join(x for x in c[1:] if x)}); i+=1; continue
        if not first and len(c)>1:
            # an indented table: gather consecutive indented rows
            tbl=[]
            while i < len(rows) and not rows[i][0] and len(rows[i])>1:
                tbl.append(rows[i][1:]); i+=1
            if tbl:
                kind = "grid" if all(len(t)>3 for t in tbl) and len(tbl)>2 and not tbl[0][0] else "table"
                sheet["blocks"].append({"kind":kind,"rows":tbl})
            continue
        if len(c)==1:
            kind = "heading" if len(first)<48 and not first.endswith(".") else "text"
            sheet["blocks"].append({"kind":kind,"text":first}); i+=1; continue
        sheet["blocks"].append({"kind":"text","text":" ".join(x for x in c if x)}); i+=1
    for s in sheet["samples"]:
        p = s.get("product")
        if p: sheet["outputs"].append({"construct": p, "label": s.get("label","")})
    sheet["id"] = f"{name.lower().replace(' ','_')}"
    if sheet["operation"].lower().startswith("dilution"):
        d = dilution_of(rows, descriptions or {})
        if d:
            sheet["dilution"] = d
            # The prose walk is REPLACED, not kept beside the structure. Keeping both would
            # print the calculation and then a page telling the student to do it by hand.
            sheet["blocks"] = []
            print(f"     {name}: {len(d['targets'])} oligo(s) to {d['target_uM']:g} uM "
                  f"from a {d['stock_uM']:g} uM stock")
        else:
            print(f"     ! {name}: looks like a dilution step but names no oligos — "
                  f"left as prose. Check the tab.")
    if sheet.get("applies") is False:
        print(f"     - {name}: marked False in the workbook — carried, but not a step of this "
              f"experiment ({sheet['title'][:50]!r})")
    return sheet

# WHICH TABS ARE STEPS — asked of the workbook, never of a list in here.
#
# This was a hardcoded ordered list of fourteen names, fitted to SLIP4-7. It worked for exactly
# that workbook. SLIP4-8 assembles by `Gibson`, SLIP4_libraries spells it `Golden Gate` with a
# space and adds `Zymo2` and `Seq Analysis`, SLIP4_singles adds `PCR2`. Every one of those names
# was absent from the list, so each would have been dropped — and the packet would have looked
# complete, with the assembly step simply not in it. That is absence collapsing into a clean
# result, which is the failure this whole toolchain is built to refuse.
#
# So the rule is inverted: a tab is a step unless it is one of the known reference tabs. The
# workbook's own tab order IS the protocol order — a new operation therefore arrives on its own,
# and a new *reference* tab shows up as an unexpected sheet, which is the failure that is loud.
REFERENCE_TABS = {"sequences", "construction", "calculations", "inventory", "stock materials",
                  "pl prep", "instructions"}

# DILUTION IS AN OPERATION, NOT A PAGE OF PROSE.
#
# JCA, 2026-09-10, looking at a rendered SLIP5 Dilution sheet that was six lines of
# "oligo oGho23 Removal of BsaI site…": *"That's a pretty crappy labsheet. Dilution should be
# one of the standard types defined in labsheet, with oligo names and target concentrations.
# The compiler would look up the description and insert it. It should be presented as the
# calculated type — so student puts in the nmol number, and it says conc to make."*
#
# The two workbooks state it two different ways and mean the same three things: a TARGET
# concentration, a LIST of oligos, and what each one is FOR. SLIP5 writes `oligo <name>
# <description>` under "Resuspend and make 10 uM dilutions of:". SLIP4-7 writes a conditional
# walk — check for working stocks, check for IDT stocks, resuspend, dilute — with a
# Construct/Box/Well/Concentration table at each branch. Reading both into one structure is
# what makes the rendered page a calculation instead of a transcription.
TARGET = re.compile(r"(\d+(?:\.\d+)?)\s*uM\b", re.I)
HEADERWORDS = {"construct", "label", "side-label", "box", "well", "concentration",
               "final concentration", "nmol", "ddh2o to add", "name", "sequence"}

def dilution_of(rows, descriptions):
    """{stock_uM, target_uM, targets:[{oligo, description, box, well}]} or None."""
    target = stock = None
    found, order, cols = {}, [], {}

    def add(name, desc="", box="", well=""):
        if not name or DNAish(name): return
        if name not in found:
            found[name] = {"oligo": name, "description": desc, "box": box, "well": well}
            order.append(name)
        else:
            for k, v in (("description", desc), ("box", box), ("well", well)):
                if v and not found[name][k]: found[name][k] = v

    for c in rows:
        line = " ".join(x for x in c if x)
        low = line.lower()
        m = TARGET.search(line)
        if m:
            # "make 10 uM dilutions" / "you will need 10 uM stocks" name the TARGET;
            # "these 100 uM stocks" names the IDT stock the target is made from.
            if "stock" in low and float(m.group(1)) >= 50: stock = float(m.group(1))
            elif target is None: target = float(m.group(1))
        if c and c[0].lower() == "oligo" and len(c) > 1:
            add(c[1], c[2] if len(c) > 2 else ""); continue
        # A HEADER ROW IS A HEADER ROW, and it is also where the columns are named. Reading
        # `| nmol | ddH2O to Add | Box | Well |` as data invented an oligo called "nmol" whose
        # box was "ddH2O to Add"; taking Box and Well from fixed offsets happened to be right
        # on one table and wrong on the next, where the same tab puts Label and Side-Label in
        # between. Both are fixed by using the header for what it is.
        if any(x.lower() in HEADERWORDS for x in c):
            hdr = {x.lower(): i for i, x in enumerate(c)}
            # The construct column is often unlabelled — the blank cell left of `Label`.
            cols["name"] = (hdr.get("construct") if "construct" in hdr else
                            next((i for i, x in enumerate(c) if i and not x), None))
            cols["box"], cols["well"] = hdr.get("box"), hdr.get("well")
            continue
        if cols.get("name") is not None and len(c) > cols["name"] and c[cols["name"]]:
            g = lambda k: (c[cols[k]] if cols.get(k) is not None and len(c) > cols[k] else "")
            add(c[cols["name"]], "", g("box"), g("well"))
    if not order: return None
    for t in order:
        # THE COMPILER LOOKS UP THE DESCRIPTION. An oligo named on the dilution page and
        # described on the sequences page is one fact in two places; the student should not
        # have to hold the name in their head while they turn to another tab.
        if not found[t]["description"]: found[t]["description"] = descriptions.get(t, "")
    return {"stock_uM": stock or 100.0, "target_uM": target or 10.0,
            "targets": [found[t] for t in order]}


def DNAish(s):
    return bool(re.fullmatch(r"[ACGTacgtNn]{8,}", s or ""))


def descriptions_in(wb):
    """{name: what it is for} from whichever tab states it — usually `sequences`."""
    out = {}
    for ws in wb.worksheets:
        if ws.title.strip().lower() in ("calculations",): continue
        for c in cells(ws):
            v = [x for x in c if x]
            if len(v) >= 3 and DNAish(v[1]) and not DNAish(v[2]) and len(v[0]) < 40:
                out.setdefault(v[0], v[2])
    return out


def experiment_of(title):
    """The experiment a tab claims to belong to, or None if it does not say."""
    m = re.search(r"for\s+Experiment\s+(\S+)", title or "", re.I)
    return m.group(1).strip() if m else None


def step_tabs(wb):
    steps   = [t for t in wb.sheetnames if t.strip().lower() not in REFERENCE_TABS]
    skipped = [t for t in wb.sheetnames if t.strip().lower() in REFERENCE_TABS]
    return steps, skipped


def main(src, out):
    wb = openpyxl.load_workbook(src, data_only=True)
    steps, skipped = step_tabs(wb)
    descs = descriptions_in(wb)
    packet = {"id": os.path.basename(src).replace(".xlsx",""),
              "metadata": {"title": None, "experiment": None,
                           "source": os.path.basename(src)}, "sheets": []}
    for t in steps:
        if t not in wb.sheetnames: continue
        s = convert_tab(wb[t], t, descs)
        if s: packet["sheets"].append(s)
    if packet["sheets"]:
        title = packet["sheets"][0]["title"]
        packet["metadata"]["experiment"] = title.split("Experiment",1)[-1].strip() if "Experiment" in title else ""
        packet["metadata"]["title"] = f"LabSheets — {packet['metadata']['experiment']}"
    # DOES EVERY TAB AGREE ABOUT WHICH EXPERIMENT THIS IS?
    #
    # These workbooks are made by copying last one and editing it, so a tab that nobody had to
    # change keeps the previous experiment's name in its title. SLIP4-7 carries four tabs
    # headed SLIP4-8 and one headed SLIP4-1. It is cosmetic right up until a student prints
    # the packet, works from the Miniprep page, and files the result under the wrong
    # experiment — and it is invisible to anyone reading one tab at a time, which is how a
    # workbook is always read. Comparing them costs nothing and can only be done here, where
    # all of the tabs are in view at once.
    ours = packet["metadata"].get("experiment")
    # A tab already reported as untitled is not ALSO a naming mismatch — one defect, one line.
    odd = [(s2["id"], exp) for s2 in packet["sheets"]
           for exp in [experiment_of(s2["title"])] if exp and ours and exp != ours]
    if odd:
        print(f"     ! {len(odd)} tab(s) name a different experiment than {ours}:")
        for sid, exp in odd: print(f"         {sid:<16} says {exp}")

    json.dump(packet, open(out,"w"), indent=2)
    print(f"  wrote {out}: {len(packet['sheets'])} sheet(s)")
    print(f"     reference tabs not converted: {', '.join(skipped) if skipped else '(none)'}")
    for s in packet["sheets"]:
        print(f"     {s['title'][:44]:46} inputs={len(s['inputs'])} samples={len(s['samples'])} "
              f"recipe={len(s['recipe']['components']) if s['recipe'] else 0} notes={len(s['notes'])} blocks={len(s['blocks'])}")

# Importable, so the conversion rules can be tested without a workbook on disk. It ran on
# import before, which meant `import xlsx_to_labpacket` tried to convert sys.argv[1].
if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
