"""Every cell a student is asked to fill in can be read back.

**A SHADED CELL WITH NO SLUG IS A QUESTION NOBODY CAN HEAR THE ANSWER TO.** The workbook shades
what it wants filled in, and `render/read-returned.py` reads answers back by slug from the record
tab. A cell in the first set and not the second is asked for, written into, and then unreachable.

Found 2026-09-13 by comparing the two sets on a real workbook: **46 of 91 shaded cells had no
slug.** Two kinds, and the second is the worse:

  `Your notes` — four lines on every sheet, under a sentence reading *"Anything that happened that
  the plan did not say. This comes back with the file and becomes part of the record."* A promise
  the file could not keep: a student's account of what actually happened, which is the part no plan
  anticipated and the part most worth having, was written down and reachable by nobody.

  `The single clone you are most confident about` on the analysis sheet — which its own design note
  calls the answer *"every session after this one"* depends on. Only the Samples table registered
  its entry cells; a block table did not, and that answer lives in a block table.
"""
import os
import re
import subprocess
import sys
import tempfile
import unittest

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
C6 = os.path.dirname(os.path.dirname(HERE))
FIXTURE = os.path.join(C6, "test", "fixtures", "golden")
RENDER = os.path.join(C6, "src", "labplanner", "render", "labpacket-to-xlsx.py")
ENTRY = "FFF6C8"
REF = re.compile(r"'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)")


def _workbook(tmp):
    got = subprocess.run(["node", os.path.join(C6, "bin", "c6-packet"), FIXTURE,
                          "--inventory", os.path.join(FIXTURE, "inventory.txt")],
                         capture_output=True, text=True)
    assert got.returncode == 0, got.stderr[-400:]
    pj = os.path.join(tmp, "p.json")
    open(pj, "w").write(got.stdout)
    out = os.path.join(tmp, "w.xlsx")
    made = subprocess.run([sys.executable, RENDER, pj, out], capture_output=True, text=True)
    assert made.returncode == 0, made.stderr[-800:]
    return openpyxl.load_workbook(out)


class TestEveryQuestionIsAnswerable(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        cls.wb = _workbook(cls._tmp.name)
        cls.captured = set()
        for row in cls.wb["record"].iter_rows(min_row=2, max_col=2, values_only=True):
            m = REF.search(str((row + (None, None))[1] or ""))
            if m:
                cls.captured.add((m.group(1), f"{m.group(2)}{m.group(3)}"))

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def _shaded(self):
        for ws in self.wb.worksheets:
            if ws.title == "record":
                continue
            for row in ws.iter_rows():
                for c in row:
                    f = c.fill
                    if f and f.fgColor and str(f.fgColor.rgb or "")[-6:] == ENTRY:
                        yield ws, c

    def test_the_workbook_asks_for_something(self):
        """A workbook with no entry cells would pass everything below vacuously."""
        self.assertGreater(len(list(self._shaded())), 20)

    def test_every_shaded_cell_has_a_slug(self):
        missing = [f"{ws.title}!{c.coordinate}" for ws, c in self._shaded()
                   if (ws.title, c.coordinate) not in self.captured]
        self.assertEqual(missing, [], f"{len(missing)} asked-for cell(s) cannot be read back")

    def test_the_notes_block_is_readable(self):
        """Its own sentence says it becomes part of the record."""
        slugs = [s for s, in [(r[0],) for r in
                              self.wb["record"].iter_rows(min_row=2, max_col=1, values_only=True)]
                 if s and ".notes." in str(s)]
        self.assertGreaterEqual(len(slugs), 4)

    def test_every_slug_points_at_a_cell_that_exists(self):
        """The other direction: a slug naming a cell that is not there reads back as blank."""
        for row in self.wb["record"].iter_rows(min_row=2, max_col=2, values_only=True):
            slug, formula = (row + (None, None))[:2]
            if not slug:
                continue
            m = REF.search(str(formula or ""))
            self.assertIsNotNone(m, f"{slug}: {formula!r} is not a reference")
            self.assertIn(m.group(1), self.wb.sheetnames, f"{slug} points at a missing sheet")

    def test_no_slug_is_used_twice(self):
        """Two cells under one slug is one of them being silently unreadable."""
        seen = {}
        for row in self.wb["record"].iter_rows(min_row=2, max_col=2, values_only=True):
            slug = (row + (None,))[0]
            if not slug:
                continue
            self.assertNotIn(slug, seen, f"{slug} appears twice")
            seen[slug] = True


if __name__ == "__main__":
    unittest.main()
