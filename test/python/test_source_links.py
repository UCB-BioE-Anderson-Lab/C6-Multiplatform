"""A fact written down on one sheet is referenced by the next, not asked for again.

JCA, 2026-09-12, looking at a blank Box and Well beside "the 10 µM working stock you made in
session 1": *"use a formula to pull that info from the previous page."*

**Asking twice for one fact is how the second answer comes back different**, and the second ask is
the one nobody fills in. The dilution session records where each working stock went; the PCR sheet
that uses it points at those cells and follows whatever the student typed.
"""
import sys
import re
import os
import importlib.util, json, os, subprocess, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, "..", ".."))
MOD = os.path.join(ROOT, "src", "labplanner", "render", "labpacket-to-xlsx.py")
spec = importlib.util.spec_from_file_location("lp2x", MOD)
lp = importlib.util.module_from_spec(spec); spec.loader.exec_module(lp)

CF = "PCR\tbo1\tbo2\tpSRC\tfrag1\nGoldenGate\tfrag1\tbackbone\tBsaI\tpTEST\n"
OLIGOS = "bo1\tACGTACGTACGTACGTACGT\t25nm\tSTD\nbo2\tTGCATGCATGCATGCATGCA\t25nm\tSTD\n"


def _packet():
    d = tempfile.mkdtemp(prefix="links-")
    open(os.path.join(d, "Construction of pTEST.txt"), "w").write(CF)
    open(os.path.join(d, "p_oligos.txt"), "w").write(OLIGOS)
    out = subprocess.run(["node", os.path.join(ROOT, "bin", "c6-packet"), d, "--clone-only"],
                         capture_output=True, text=True)
    assert out.returncode == 0, out.stderr[:400]
    return json.loads(out.stdout)


PACKET = _packet()
PCR = next(s for s in PACKET["sheets"] if "pcr" in s["metadata"]["operations"])


def test_the_pcr_sheet_points_at_the_dilution_sheet_rather_than_asking():
    rows = {r["what"]: r for r in PCR["sources"]}
    for o in ("bo1", "bo2"):
        assert rows[o].get("link"), rows[o]
        assert rows[o]["link"]["box"].endswith(f".working.{o}.box"), rows[o]["link"]
        # Not also an ask: two yellow cells for one fact is the thing being removed.
        assert not rows[o].get("unlocated") and not rows[o].get("askWell"), rows[o]


def test_the_slug_it_points_at_is_one_the_dilution_sheet_actually_writes():
    """A reference to a slug nobody records is a blank cell and a silent lie about linkage."""
    d = next(s for s in PACKET["sheets"] if s.get("dilution"))
    slug = d["dilution"]["slug"]
    for r in PCR["sources"]:
        if r.get("link"):
            assert r["link"]["box"].startswith(f"{slug}."), (r["link"], slug)


def test_the_reference_is_guarded_so_an_unfilled_cell_is_blank_and_not_zero():
    """A bare =E18 against an empty cell displays 0. A student reading "0" in a Box column either
    goes looking for box zero or decides the sheet is broken."""
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "PCR"
    record = [("s1-dilution.working.bo1.box", "'Oligo dilutions'!E18"),
              ("s1-dilution.working.bo1.well", "'Oligo dilutions'!F18")]
    sheet = {"id": "s2-pcr", "inputs": [
        {"what": "bo1", "box": "", "well": "", "made": True,
         "link": {"box": "s1-dilution.working.bo1.box",
                  "well": "s1-dilution.working.bo1.well"},
         "note": "made in session 1"}]}
    lp.write_sources(ws, 1, sheet, record)
    got = [c.value for row in ws.iter_rows(values_only=False) for c in row if c.value]
    formulas = [v for v in got if str(v).startswith("=")]
    assert len(formulas) == 2, got
    for f in formulas:
        assert f.startswith("=IF(") and '="",""' in f, f


def test_an_unresolvable_link_leaves_the_cells_empty():
    """Rather than a #REF!, which on a printed labsheet reads as a system fault."""
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "PCR"
    sheet = {"id": "s2-pcr", "inputs": [
        {"what": "bo1", "link": {"box": "nobody.writes.this.box",
                                 "well": "nobody.writes.this.well"}, "note": "x"}]}
    lp.write_sources(ws, 1, sheet, [])
    assert not [c.value for row in ws.iter_rows() for c in row
                if c.value and str(c.value).startswith("=")]


def test_a_box_that_does_not_track_wells_asks_for_nothing():
    """JCA, 2026-09-12: *"It is not worthwhile to speak of the location of pJ01. It is often used,
    and it moves around in that box as a result."* The box is printed and no cell is yellow — a
    question whose answer goes stale immediately trains people to skip the ones that do not."""
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "PCR"
    sheet = {"id": "s1-pcr", "inputs": [
        {"what": "pJ01", "box": "Pink Training", "well": "", "note": "Miniprep DNA in Pink Training."},
        {"what": "pOTHER", "box": "Cheese1", "well": "", "askWell": True, "note": "well not recorded"}]}
    record = []
    lp.write_sources(ws, 1, sheet, record)
    # Only the one that is genuinely unknown gets a slug and a cell to fill.
    assert [s for s, _ in record] == ["s1-pcr.source.pOTHER.well"], record


def test_the_clone_is_asked_for_once_and_computed_from():
    """JCA, 2026-09-12: *"It is unknown when we write the labsheet which clone is being taken into
    this, so it is appropriately written up as just pBET8. When the student goes to do this, they
    will type in which clone(s) they are applying it to, so the spreadsheet should give a place for
    them to put in this info, and then you can refer to the full actual clone name (pBET8-A)
    calculated from the supplied clone designation."*

    One yellow cell for the letter, and every reference to that construct becomes a formula.
    Asking for the whole name would be asking somebody to retype `pBET8-` and get it wrong;
    printing one of the candidates would be choosing for them.
    """
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Electroporation"
    record = []
    sheet = {"id": "s9-retransform", "inputs": [
        {"what": "pBET8", "askClone": "pBET8", "box": "", "well": "", "made": True,
         "note": "made by the transform step"}]}
    lp.write_sources(ws, 1, sheet, record)
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any("Which clone of pBET8" in c for c in cells), cells
    formula = [c for c in cells if c.startswith("=")]
    assert len(formula) == 1, cells
    # Guarded, so it is blank until they answer rather than showing a dangling dash.
    assert formula[0].startswith("=IF(") and '"pBET8-"&' in formula[0], formula
    assert [s for s, _ in record] == ["s9-retransform.clone.pBET8"], record


def test_a_step_that_makes_the_clones_is_not_asked_which_clone():
    """The pick that FEEDS the analysis consumes the same construct and is where the clones come
    FROM. Asking it which clone is being used is backwards, and it did."""
    packet = _packet()
    for sh in packet["sheets"]:
        for i in sh.get("inputs", []):
            if i.get("askClone"):
                ops = sh["metadata"]["operations"]
                assert "pick" not in ops, (sh["id"], ops)


if __name__ == "__main__":
    fails = []
    # **A TEST DEFINED AFTER THIS BLOCK IS NOT RUN, AND THE FILE STILL PRINTS "passed".** Six new
    # tests were appended to this file on 2026-09-13 below the runner; `globals()` had not seen them
    # yet, so the suite collected the old five, reported green, and the new ones had never executed.
    # That is the shape `CLAUDE.md` keeps naming — *"a test nothing runs is the same failure one
    # layer down"* — and a hand-rolled runner cannot see it from the inside without being asked to.
    _src = open(os.path.abspath(__file__), encoding="utf-8").read()
    _declared = set(re.findall(r"^def (test_\w+)", _src, re.M))
    _collected = {n for n in globals() if n.startswith("test_") and callable(globals()[n])}
    if _declared - _collected:
        print(f"  MISSED  {len(_declared - _collected)} test(s) defined below the runner and never "
              f"run: {', '.join(sorted(_declared - _collected))}")
        sys.exit(1)
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            # CATCH EVERYTHING, NOT ONLY AssertionError. A KeyError raised by one test used to
            # escape this loop, so every test after it never ran AND the summary line never
            # printed — leaving output that greps as clean. One renamed field hid a whole file.
            try: fn(); print(f"  ok    {name}")
            except AssertionError as e:
                fails.append(name); print(f"  FAIL  {name}: {str(e)[:300]}")
            except Exception as e:
                fails.append(name)
                print(f"  ERROR {name}: {type(e).__name__}: {str(e)[:300]}")
    print(f"\n{'FAILED' if fails else 'passed'}: {len(fails)} failure(s)")
    sys.exit(1 if fails else 0)
