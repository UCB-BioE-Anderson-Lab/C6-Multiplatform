"""A long sheet is several pages, not one page shrunk until nobody can read it.

`fitToHeight = 1` was a promise the page setup could not keep. It scaled a 68-row assay tab to
whatever percentage made it fit — at a bench, under gloves — while the footer said "Page 1 of 1"
and the pipeline printed a warning telling somebody to "cut what it says". Two of the four
overflowing sheets are long precisely because they transclude a protocol that has no cheatsheet,
which is the content JCA asked to be included: *"It's the protocols that don't have cheatsheets
that need to be included on the labsheets."*

So the renderer paginates, and the breaks land between sections rather than through a table.
"""
import json
import os
import subprocess
import sys
import tempfile
import unittest

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.normpath(os.path.join(HERE, "..", ".."))
RENDER = os.path.join(REPO, "src", "labplanner", "render", "labpacket-to-xlsx.py")
FIXTURE = os.path.join(REPO, "test", "fixtures", "golden")


def _render(tmp):
    got = subprocess.run(
        ["node", os.path.join(REPO, "bin", "c6-packet"), FIXTURE,
         "--inventory", os.path.join(FIXTURE, "inventory.txt")],
        capture_output=True, text=True)
    assert got.returncode == 0, got.stderr[:400]
    pj = os.path.join(tmp, "packet.json")
    with open(pj, "w") as f:
        f.write(got.stdout)
    out = os.path.join(tmp, "out.xlsx")
    made = subprocess.run([sys.executable, RENDER, pj, out],
                          capture_output=True, text=True)
    assert made.returncode == 0, made.stderr[:800]
    return openpyxl.load_workbook(out), json.loads(got.stdout)


class TestPagination(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        cls.wb, cls.packet = _render(cls._tmp.name)

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def test_no_sheet_is_squashed_to_fit_one_page(self):
        for ws in self.wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            self.assertEqual(ws.page_setup.fitToHeight, 0,
                             f"{ws.title} is being scaled down to one page")
            self.assertEqual(ws.page_setup.fitToWidth, 1,
                             f"{ws.title} is not fit to the page width")

    def test_a_long_sheet_gets_at_least_one_break(self):
        long_ones = [ws for ws in self.wb.worksheets
                     if ws.sheet_state == "visible" and ws.max_row > 44]
        self.assertTrue(long_ones, "the fixture no longer has a sheet over one page")
        for ws in long_ones:
            self.assertTrue(list(ws.row_breaks.brk),
                            f"{ws.title} is {ws.max_row} rows and has no page break")

    def test_every_break_lands_on_a_section_heading(self):
        """A break through the middle of a table strands its header on the page before."""
        navy = "1F3864"
        shaded = "DCE6F1"
        for ws in self.wb.worksheets:
            for b in ws.row_breaks.brk:
                c = ws.cell(row=b.id + 1, column=1)
                self.assertTrue(c.value, f"{ws.title}: break lands above an empty row")
                self.assertTrue(c.font and c.font.bold and c.font.color
                                and str(c.font.color.rgb)[-6:] == navy,
                                f"{ws.title}: break above {c.value!r}, not a section heading")
                fill = c.fill
                if fill is not None and fill.fgColor is not None:
                    self.assertNotEqual(str(fill.fgColor.rgb)[-6:], shaded,
                                        f"{ws.title}: break above a table's column header")

    def test_the_footer_page_count_is_now_true(self):
        """"Page &P of &N" was printed on every sheet while the setup forced N to 1."""
        for ws in self.wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            self.assertIn("&P", ws.oddFooter.right.text or "")


if __name__ == "__main__":
    unittest.main()
